import io
import os
import re
import base64
import logging
from typing import Tuple, List, Dict, Any
from pypdf import PdfReader
from docx import Document

logger = logging.getLogger(__name__)


def extract_scanned_documents_with_gemini(images_data: List[Tuple[bytes, str]]) -> str:
    """
    Extract structured credentials, certifications, awards, and trade details from
    scanned PDF image pages using Gemini Vision AI.
    Processes images in batches with strict timeout and fallback handling.
    """
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key or not images_data:
        return ""

    models_to_try = [
        "gemini-3.5-flash-lite",
        "gemini-3.1-flash-lite",
        "gemini-3.8-flash",
    ]

    prompt = (
        "The following images are attached certificates, diplomas, training credentials, "
        "service letters, or awards from a candidate's CV.\n"
        "For EACH document/certificate shown, extract all factual details:\n"
        "- Document / Certification / Award Title\n"
        "- Issuing Organization / Institution / Authority\n"
        "- Candidate / Recipient Name\n"
        "- Issue Date / Year / Period\n"
        "- Credential ID or Reference Number (if any)\n"
        "- Grade, NVQ Level, Score, or Distinction (if any)\n"
        "- Key Skills, Trade Competencies, or Responsibilities mentioned\n\n"
        "Format each certificate clearly as a distinct item so that the AI resume parser "
        "can accurately include it in the candidate's Certifications and Qualifications list."
    )

    extracted_sections = []
    # Process in batches of up to 6 images to avoid payload timeouts
    batch_size = 6
    for b_idx in range(0, len(images_data), batch_size):
        batch = images_data[b_idx:b_idx + batch_size]
        batch_text = ""

        # 1. Try google.genai Client
        try:
            from google import genai
            from google.genai import types

            client = genai.Client(
                api_key=api_key,
                http_options=types.HttpOptions(timeout=25000)
            )

            parts = [prompt]
            for img_bytes, img_name in batch:
                mime = "image/jpeg"
                img_name_lower = (img_name or "").lower()
                if img_name_lower.endswith(".png"):
                    mime = "image/png"
                elif img_name_lower.endswith(".webp"):
                    mime = "image/webp"
                parts.append(types.Part.from_bytes(data=img_bytes, mime_type=mime))

            for model_name in models_to_try:
                try:
                    response = client.models.generate_content(
                        model=model_name,
                        contents=parts
                    )
                    if response and response.text and response.text.strip():
                        batch_text = response.text.strip()
                        break
                except Exception as me:
                    logger.warning("[Vision OCR] Model %s failed: %s", model_name, me)
                    continue
        except Exception as e:
            logger.warning("[Vision OCR] google.genai client init failed: %s. Trying REST API...", e)

        # 2. REST API fallback if client failed
        if not batch_text:
            try:
                import requests
                for model_name in models_to_try:
                    try:
                        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
                        payload_parts: List[Dict[str, Any]] = [{"text": prompt}]
                        for img_bytes, img_name in batch:
                            mime = "image/jpeg"
                            img_name_lower = (img_name or "").lower()
                            if img_name_lower.endswith(".png"):
                                mime = "image/png"
                            elif img_name_lower.endswith(".webp"):
                                mime = "image/webp"
                            b64 = base64.b64encode(img_bytes).decode("utf-8")
                            payload_parts.append({
                                "inline_data": {
                                    "mime_type": mime,
                                    "data": b64
                                }
                            })

                        req_payload = {
                            "contents": [{"parts": payload_parts}]
                        }
                        r = requests.post(url, json=req_payload, timeout=25)
                        if r.status_code == 200:
                            data = r.json()
                            candidates = data.get("candidates", [])
                            if candidates:
                                content_parts = candidates[0].get("content", {}).get("parts", [])
                                if content_parts:
                                    batch_text = content_parts[0].get("text", "").strip()
                                    if batch_text:
                                        break
                    except Exception as re_err:
                        logger.warning("[Vision OCR REST] %s failed: %s", model_name, re_err)
                        continue
            except Exception as rest_ex:
                logger.error("[Vision OCR REST] Top level error: %s", rest_ex)

        if batch_text:
            extracted_sections.append(batch_text)

    return "\n\n".join(extracted_sections).strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract plain text from a PDF file stream using pypdf.
    If scanned certificate or image pages are detected (pages with little/no text but images),
    Gemini Vision AI is automatically invoked to extract credentials and certificates.
    """
    reader = PdfReader(io.BytesIO(file_bytes))
    extracted_pages = []
    scanned_images = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        has_text = bool(text and len(text.strip()) >= 30)

        if has_text:
            extracted_pages.append(text.strip())
        else:
            # Check for scanned page images (certificates, awards, credentials)
            try:
                if hasattr(page, "images") and page.images:
                    for img in page.images:
                        # Only take substantive images (> 5KB) to filter out icons/spacers
                        if len(img.data) > 5000:
                            scanned_images.append((img.data, img.name))
                            break
            except Exception as img_err:
                logger.warning("[PDF Extract] Could not inspect images on page %d: %s", i + 1, img_err)

    # If scanned document pages were detected, perform Vision AI OCR
    if scanned_images:
        logger.info("[PDF Extract] Detected %d scanned certificate/document pages. Invoking Vision AI OCR...", len(scanned_images))
        try:
            ocr_text = extract_scanned_documents_with_gemini(scanned_images)
            if ocr_text:
                extracted_pages.append(
                    "\n\n=== CERTIFICATIONS & VERIFIED CREDENTIALS (EXTRACTED VIA VISION AI) ===\n" + ocr_text
                )
        except Exception as ocr_ex:
            logger.error("[PDF Extract] Vision OCR failed gracefully: %s", ocr_ex)

    full_text = "\n\n".join(extracted_pages)
    return clean_extracted_text(full_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from a DOCX file stream using python-docx."""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # Also extract text inside tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_text:
                paragraphs.append(row_text)
                
    full_text = "\n".join(paragraphs)
    return clean_extracted_text(full_text)


def clean_extracted_text(text: str) -> str:
    """Normalize whitespace, remove non-printable chars, and standardize bullets."""
    if not text:
        return ""
    # Strip unprintable or odd control chars like \x7f and unicode squares
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    # Replace weird unicode bullets with standard dashes
    text = re.sub(r'[\u2022\u2023\u25E6\u2043\u2219\u25A0\u25AA\u25AB\u25CF\u25CB\u25BA\u25B6\uF0B7■▪●•]', '\n- ', text)
    # Standardize multiple newlines
    text = re.sub(r'\n{3,}', '\n\n', text)
    # Remove excessive horizontal whitespace
    text = re.sub(r'[ \t]{2,}', ' ', text)
    return text.strip()


def detect_sections(text: str) -> List[str]:
    """Detect presence of common standard resume sections."""
    detected = []
    text_lower = text.lower()
    
    section_patterns = [
        ("Summary / Objective", r"(summary|objective|profile|about me)"),
        ("Work Experience", r"(experience|employment|work history|career history|professional background|work experience)"),
        ("Education", r"(education|academic background|degrees|university|college|bsc|hnd)"),
        ("Skills & Technologies", r"(skills|technical skills|competencies|technologies|programming languages|frameworks)"),
        ("Projects & Works", r"(projects|portfolio|personal projects|key projects|projects and my works)"),
        ("Certifications & Awards", r"(certifications|certificates|licenses|credentials|awards)")
    ]
    
    for label, pattern in section_patterns:
        if re.search(pattern, text_lower):
            detected.append(label)
            
    return detected


NAME_BLACKLIST = set([
    'university', 'college', 'institute', 'academy', 'school', 'polytechnic', 'campus', 'faculty', 'department',
    'pvt', 'ltd', 'inc', 'llc', 'corp', 'corporation', 'company', 'systems', 'technologies', 'services', 'solutions',
    'road', 'street', 'avenue', 'lane', 'drive', 'boulevard', 'city', 'state', 'country', 'sri lanka', 'colombo', 'beliatta',
    'resume', 'curriculum', 'vitae', 'cv', 'page', 'objective', 'summary', 'profile', 'education', 'experience', 'skills',
    'personal', 'contact', 'nationality', 'gender', 'dob', 'date', 'birth', 'passport', 'nic', 'qualifications', 'references',
    'projects', 'works', 'portfolio', 'languages', 'frameworks', 'database', 'certifications', 'awards', 'links', 'socials',
    'engineer', 'developer', 'manager', 'analyst', 'designer', 'architect', 'consultant', 'technician', 'electrician',
    'painter', 'specialist', 'executive', 'officer', 'director', 'intern', 'associate', 'lead', 'founder', 'freelance',
    'bsc', 'bachelor', 'master', 'msc', 'phd', 'diploma', 'hnd', 'nvq', 'degree', 'cardiff', 'metropolitan', 'nibm',
    'january', 'february', 'march', 'april', 'may', 'june', 'july', 'august', 'september', 'october', 'november', 'december',
    'present', 'year', 'years', 'month', 'months', 'repository', 'website', 'application', 'applications', 'tools', 'database',
    'surface', 'preparation', 'sanding', 'coating', 'primer', 'polishing', 'diagnostics', 'inspection', 'wiring', 'repairs',
    'assembly', 'welding', 'fabrication', 'finishes', 'harnesses', 'lighting', 'alternator', 'testing', 'sensor', 'charging',
    'technical', 'competencies', 'responsibilities', 'duties', 'achievements', 'hardware', 'troubleshooting', 'safety'
])


def extract_candidate_name(text: str, email: str = "") -> str:
    """Intelligently extract the candidate's actual full name from text with confidence scoring."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return "Candidate Name"

    email_user = email.split('@')[0].lower() if email else ''
    candidates = []

    for i, line in enumerate(lines):
        # Ignore lines with obvious non-name symbols
        if any(c in line for c in [':', '|', '@', '/', '\\', '=', '+', ';', '#', '(', ')', '{', '}', '[', ']']):
            continue
        if re.search(r'\d', line):
            continue

        cleaned = re.sub(r'[^a-zA-Z\s]', '', line).strip()
        words = cleaned.split()
        if not (2 <= len(words) <= 6):
            continue
        # Each word must start with an uppercase letter and consist only of letters
        if not all(w.isalpha() and w[0].isupper() for w in words):
            continue

        line_lower = line.lower()
        # Strictly reject if any blacklisted keyword exists as a word or substring
        if any(re.search(rf'\b{re.escape(b)}\b', line_lower) for b in NAME_BLACKLIST):
            continue
        if any(b in line_lower for b in ['university', 'college', 'institute', 'school', 'cardiff', 'nibm', 'pvt', 'ltd', 'road', 'street']):
            continue

        score = 10
        # Positional priority - first 3 lines are overwhelmingly candidate names
        if i == 0:
            score += 60
        elif i == 1:
            score += 45
        elif i < 5:
            score += 25
        elif i < 15:
            score += 15
        elif i < 80:
            score += 5

        # Adjacent title boost (e.g. line right after is 'Automotive Technician' or 'Software Engineer')
        if i + 1 < len(lines):
            next_line = lines[i+1].lower()
            if any(t in next_line for t in ['engineer', 'developer', 'manager', 'specialist', 'consultant', 'technician', 'lead', 'architect', 'painter', 'electrician', 'mechanic', 'officer']):
                score += 80

        # Email match boost (e.g. 'isuranga' in 'lghisuranga@gmail.com')
        for w in words:
            w_low = w.lower()
            if len(w_low) >= 3 and (w_low in email_user or email_user in w_low or (len(w_low) >= 4 and w_low[:4] in email_user)):
                score += 90

        candidates.append((score, cleaned, i))

    if candidates:
        candidates.sort(key=lambda x: x[0], reverse=True)
        return candidates[0][1]

    # Email username fallback
    if email_user:
        name_parts = re.findall(r'[A-Za-z]+', email_user)
        if name_parts and len(name_parts) >= 2:
            return " ".join([p.capitalize() for p in name_parts])
        elif name_parts and len(name_parts[0]) >= 3:
            return name_parts[0].capitalize()

    return "Candidate Name"


def parse_resume_file(filename: str, file_bytes: bytes) -> Tuple[str, List[str]]:
    """Parse resume file based on extension and return (text, detected_sections)."""
    filename_lower = filename.lower()
    
    if filename_lower.endswith('.pdf'):
        text = extract_text_from_pdf(file_bytes)
    elif filename_lower.endswith(('.docx', '.doc')):
        text = extract_text_from_docx(file_bytes)
    elif filename_lower.endswith('.txt'):
        try:
            text = file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            text = file_bytes.decode('latin-1', errors='ignore')
        text = clean_extracted_text(text)
    else:
        raise ValueError(f"Unsupported file format: {filename}. Please upload a PDF, DOCX, or TXT file.")
        
    detected = detect_sections(text)
    return text, detected
